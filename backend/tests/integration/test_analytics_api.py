"""Revenue analytics reconciliation, tenant isolation, and export tests."""

from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select

from nexus.modules.analytics.models import AnalyticsExport
from nexus.modules.payments.models import PaymentIntent, PaymentRefund
from nexus.platform.background.service import JobExecutor
from nexus.platform.database.manager import DatabaseManager
from tests.integration.test_payments_api import (
    auth,
    create_intent,
    register,
    send_webhook,
    setup_confirmed_order,
    signed_payment_event,
)

pytestmark = [pytest.mark.integration, pytest.mark.anyio]
pytest_plugins = ("tests.integration.test_payments_api",)


async def complete_payment(
    client: AsyncClient, token: str, slug: str
) -> tuple[dict[str, object], dict[str, object]]:
    setup = await setup_confirmed_order(client, token, slug)
    intent = await create_intent(client, token, str(setup["order_id"]), f"intent-{slug}")
    assert intent.status_code == 201
    raw, signature = signed_payment_event(intent.json()["provider_intent_id"], "succeeded")
    assert (await send_webhook(client, raw, signature)).status_code == 200
    return setup, intent.json()


async def complete_multi_item_payment(
    client: AsyncClient, token: str, slug: str
) -> tuple[dict[str, object], dict[str, object]]:
    organization = await client.post(
        "/api/v1/organizations",
        headers=auth(token),
        json={"name": f"Analytics {slug}", "slug": slug},
    )
    branch = await client.post(
        "/api/v1/branches",
        headers=auth(token),
        json={"name": "Main branch", "slug": f"branch-{slug}"},
    )
    products = []
    for name, price in (("Primary", "10.0000"), ("Secondary", "5.0000")):
        product = await client.post(
            "/api/v1/products",
            headers=auth(token),
            json={"name": f"{name} {slug}", "price": price, "status": "active"},
        )
        assert product.status_code == 201
        products.append(product.json())
        stock = await client.post(
            "/api/v1/inventory/stocks",
            headers=auth(token),
            json={
                "branch_id": branch.json()["id"],
                "product_id": product.json()["id"],
                "low_stock_threshold": "0",
            },
        )
        seed = await client.post(
            "/api/v1/inventory/adjustments",
            headers=auth(token, f"seed-{slug}-{name}"),
            json={
                "branch_id": branch.json()["id"],
                "product_id": product.json()["id"],
                "quantity_delta": "10",
                "reason": "Analytics test seed",
            },
        )
        assert stock.status_code == seed.status_code == 201
    order = await client.post("/api/v1/orders", headers=auth(token))
    for product, quantity in zip(products, ("2", "1"), strict=True):
        item = await client.post(
            f"/api/v1/orders/{order.json()['id']}/items",
            headers=auth(token),
            json={
                "product_id": product["id"],
                "branch_id": branch.json()["id"],
                "quantity": quantity,
            },
        )
        assert item.status_code == 201
    confirmed = await client.post(
        f"/api/v1/orders/{order.json()['id']}/confirm", headers=auth(token)
    )
    assert organization.status_code == branch.status_code == order.status_code == 201
    assert confirmed.status_code == 200
    intent = await create_intent(client, token, order.json()["id"], f"intent-{slug}")
    raw, signature = signed_payment_event(intent.json()["provider_intent_id"], "succeeded")
    assert (await send_webhook(client, raw, signature)).status_code == 200
    return {
        "organization_id": organization.json()["id"],
        "product_id": products[0]["id"],
    }, intent.json()


async def test_metrics_reconcile_refunds_dates_and_tenant_isolation(
    payment_client: AsyncClient, payment_database_url: str
) -> None:
    alice_token = await register(payment_client, "analytics-alice@example.com")
    alice, intent = await complete_multi_item_payment(
        payment_client, alice_token, "analytics-alice"
    )
    refund = await payment_client.post(
        f"/api/v1/payments/intents/{intent['id']}/refunds",
        headers=auth(alice_token, "analytics-refund"),
        json={"amount": "5.0000", "reason": "Partial analytics refund"},
    )
    assert refund.status_code == 201
    simulated = await payment_client.post(
        f"/api/v1/payments/sandbox/refunds/{refund.json()['id']}/simulate",
        headers=auth(alice_token),
        json={"outcome": "succeeded"},
    )
    assert simulated.status_code == 200

    summary = await payment_client.get("/api/v1/analytics/summary", headers=auth(alice_token))
    revenue = await payment_client.get(
        "/api/v1/analytics/revenue?granularity=day", headers=auth(alice_token)
    )
    products = await payment_client.get(
        "/api/v1/analytics/products/best-selling", headers=auth(alice_token)
    )
    inventory = await payment_client.get(
        "/api/v1/analytics/inventory-valuation", headers=auth(alice_token)
    )
    branches = await payment_client.get(
        "/api/v1/analytics/revenue/branches", headers=auth(alice_token)
    )
    assert {summary.status_code, revenue.status_code, products.status_code} == {200}
    assert inventory.status_code == branches.status_code == 200
    assert summary.json() == {
        "currency": "USD",
        "total_revenue": "20.0000",
        "gross_revenue": "25.0000",
        "refunded_amount": "5.0000",
        "order_count": 1,
        "average_order_value": "20.0000",
    }
    assert len(revenue.json()["items"]) == 1
    assert revenue.json()["items"][0]["revenue"] == "20.0000"
    assert products.json()["items"][0]["product_id"] == alice["product_id"]
    assert products.json()["items"][0]["quantity_sold"] == "2.0000"
    assert inventory.json()["policy"] == "current_retail_price"
    assert inventory.json()["total_valuation"] == "125.0000"
    assert branches.json()["items"][0]["net_revenue"] == "20.0000"

    manager = DatabaseManager(payment_database_url, 2)
    async with manager.transaction() as session:
        gross = await session.scalar(
            select(func.sum(PaymentIntent.amount)).where(
                PaymentIntent.organization_id == UUID(str(alice["organization_id"])),
                PaymentIntent.status.in_(("succeeded", "refunded")),
            )
        )
        refunded = await session.scalar(
            select(func.sum(PaymentRefund.amount)).where(
                PaymentRefund.organization_id == UUID(str(alice["organization_id"])),
                PaymentRefund.status == "succeeded",
            )
        )
        assert Decimal(gross or 0) - Decimal(refunded or 0) == Decimal("20.0000")
    await manager.dispose()

    yesterday = (datetime.now(UTC).date() - timedelta(days=1)).isoformat()
    empty = await payment_client.get(
        f"/api/v1/analytics/summary?date_to={yesterday}", headers=auth(alice_token)
    )
    invalid = await payment_client.get(
        "/api/v1/analytics/summary?date_from=2026-02-02&date_to=2026-02-01",
        headers=auth(alice_token),
    )
    assert empty.json()["total_revenue"] == "0.0000"
    assert invalid.status_code == 422

    bob_token = await register(payment_client, "analytics-bob@example.com")
    await complete_payment(payment_client, bob_token, "analytics-bob")
    bob_summary = await payment_client.get(
        "/api/v1/analytics/summary", headers=auth(bob_token)
    )
    assert bob_summary.json()["total_revenue"] == "20.0000"
    assert bob_summary.json()["order_count"] == 1


async def test_csv_and_xlsx_exports_are_async_idempotent_and_downloadable(
    payment_client: AsyncClient,
    payment_database_url: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    token = await register(payment_client, "analytics-export@example.com")
    await complete_payment(payment_client, token, "analytics-export")
    monkeypatch.setattr(
        "nexus.modules.analytics.router.dispatch_job", lambda job_id: str(job_id)
    )
    manager = DatabaseManager(payment_database_url, 2)
    last_job_id: UUID | None = None

    for file_format in ("csv", "xlsx"):
        headers = auth(token, f"export-{file_format}")
        response = await payment_client.post(
            "/api/v1/analytics/exports",
            headers=headers,
            json={"file_format": file_format, "currency": "USD"},
        )
        replay = await payment_client.post(
            "/api/v1/analytics/exports",
            headers=headers,
            json={"file_format": file_format, "currency": "USD"},
        )
        assert response.status_code == replay.status_code == 202
        assert response.json()["created"] is True
        assert replay.json()["created"] is False
        assert response.json()["job_id"] == replay.json()["job_id"]
        job_id = UUID(response.json()["job_id"])
        last_job_id = job_id
        async with manager.transaction() as session:
            job = await JobExecutor(session).execute(job_id)
            assert job.status == "succeeded"
        download = await payment_client.get(
            f"/api/v1/analytics/exports/{job_id}/download", headers=auth(token)
        )
        assert download.status_code == 200
        if file_format == "csv":
            assert "total_revenue,20.0000" in download.content.decode("utf-8-sig")
        else:
            workbook = load_workbook(BytesIO(download.content), read_only=True)
            assert workbook.sheetnames == ["Summary", "Revenue by day"]

    async with manager.transaction() as session:
        assert await session.scalar(select(func.count()).select_from(AnalyticsExport)) == 2
    assert last_job_id is not None
    outsider_token = await register(payment_client, "analytics-export-outsider@example.com")
    organization = await payment_client.post(
        "/api/v1/organizations",
        headers=auth(outsider_token),
        json={"name": "Analytics outsider", "slug": "analytics-export-outsider"},
    )
    assert organization.status_code == 201
    hidden = await payment_client.get(
        f"/api/v1/analytics/exports/{last_job_id}/download",
        headers=auth(outsider_token),
    )
    assert hidden.status_code == 404
    await manager.dispose()
